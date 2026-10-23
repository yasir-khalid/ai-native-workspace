"""T3 - Capabilities: entitlement independent of agent access, pod-private artifacts, idempotency."""

from __future__ import annotations

import io

from conftest import delegation, internal_post, make_running_run
from pptx import Presentation

DECK = {"capability_id": "document.export.pptx", "idempotency_key": "deck-0001",
        "spec": {"title": "Deck", "slides": [{"title": "One", "bullets": ["a"]}]}}


async def download(user, pod: str, artifact_id: str) -> bytes:
    link = (await user.post(f"/v1/artifacts/{artifact_id}/download", pod=pod)).json()["url"]
    r = await user.http.get("/" + link.split("://", 1)[1].split("/", 1)[1])
    assert r.status_code == 200, r.text
    return r.content


async def test_ppt_denied_directly_without_entitlement(ben):
    r = await ben.post("/v1/pods/pod_beta/exports", json=DECK)
    assert r.status_code == 403 and r.json()["reason"] == "capability_not_entitled"
    caps = [c["capability_id"] for c in (await ben.get("/v1/pods/pod_beta/capabilities")).json()]
    assert "document.export.pptx" not in caps and "document.export.docx" in caps
    memo = {**DECK, "capability_id": "document.export.docx", "spec": {"title": "Memo", "sections": [{"title": "x", "bullets": ["y"]}]}}
    assert (await ben.post("/v1/pods/pod_beta/exports", json=memo)).status_code == 201


async def test_ppt_denied_through_agent_without_entitlement(ben):
    _, ev = await ben.ask("pod_beta", "Prepare a 5-slide deck on the HLXT thesis")
    pptx = [e for e in ev if e["type"] == "tool.completed" and e.get("tool") == "document.export.pptx"]
    assert pptx and pptx[0]["status"] == "denied"
    assert not [e for e in ev if e["type"] == "artifact.created"]
    assert (await ben.get("/v1/pods/pod_beta/artifacts")).json() == []


async def test_signed_delegation_cannot_widen_capabilities(svc):
    # Even a correctly signed delegation listing pptx is re-evaluated against live entitlements.
    run = await make_running_run(svc, "pod_beta", "ben.okafor")
    d = delegation("pod_beta", run, sub="ben.okafor", caps=["research.search", "document.export.pptx"])
    r = await internal_post("/v1/internal/tools/capabilities/document.export.pptx/invoke",
                            {"spec": DECK["spec"], "idempotency_key": "widen-0001"}, agent="equity-research", deleg=d)
    assert r.status_code == 403 and r.json()["reason"] == "capability_not_entitled"


async def test_agent_cannot_call_capability_it_does_not_support(svc):
    run = await make_running_run(svc, "pod_alpha", "alice.chen", agent="filing-analyst")
    d = delegation("pod_alpha", run, agent="filing-analyst", caps=["document.export.pptx"])
    r = await internal_post("/v1/internal/tools/capabilities/document.export.pptx/invoke",
                            {"spec": DECK["spec"], "idempotency_key": "unsupported-1"}, agent="filing-analyst", deleg=d)
    assert r.status_code == 403 and r.json()["reason"] == "agent_does_not_support_capability"


async def test_shared_renderer_serves_alpha_with_private_artifact(arjun, ben, bianca):
    _, ev = await arjun.ask("pod_alpha", "Prepare a 5-slide deck on our NWSC thesis")
    created = [e for e in ev if e["type"] == "artifact.created"]
    assert len(created) == 1 and ev[-1]["type"] == "run.completed"
    art = created[0]["artifact"]
    assert art["pod_id"] == "pod_alpha" and art["classification"] == "pod_confidential"
    data = await download(arjun, "pod_alpha", art["artifact_id"])
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) >= 6  # title + 5 content slides (+ sources)
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "POD CONFIDENTIAL" in texts and "BETA-PRIVATE" not in texts
    # Beta (even its PM) cannot see or fetch it.
    assert art["artifact_id"] not in str((await bianca.get("/v1/pods/pod_beta/artifacts")).json())
    r = await bianca.post(f"/v1/artifacts/{art['artifact_id']}/download", pod="pod_beta")
    assert r.status_code == 404
    listing = (await arjun.get("/v1/pods/pod_alpha/artifacts")).json()
    assert listing[0]["provenance"]["via_agent"] == "equity-research"
    assert listing[0]["provenance"]["citations"], "provenance keeps source lineage"


async def test_capability_replay_is_idempotent(arjun, svc):
    a = (await arjun.post("/v1/pods/pod_alpha/exports", json=DECK)).json()
    b = (await arjun.post("/v1/pods/pod_alpha/exports", json=DECK)).json()
    assert a["artifact_id"] == b["artifact_id"] and a["created"] and not b["created"]
    run = await make_running_run(svc, "pod_alpha", "arjun.patel")
    d = delegation("pod_alpha", run, sub="arjun.patel")
    body = {"spec": DECK["spec"], "idempotency_key": "agent-replay-1"}
    r1 = await internal_post("/v1/internal/tools/capabilities/document.export.pptx/invoke", body, agent="equity-research", deleg=d)
    r2 = await internal_post("/v1/internal/tools/capabilities/document.export.pptx/invoke", body, agent="equity-research", deleg=d)
    assert r1.json()["artifact"]["artifact_id"] == r2.json()["artifact"]["artifact_id"]
    assert r1.json()["artifact"]["classification"] == r2.json()["artifact"]["classification"] == "pod_confidential"
    assert len((await arjun.get("/v1/pods/pod_alpha/artifacts")).json()) == 2


async def test_classification_above_clearance_denied(arjun, alice):
    restricted = {**DECK, "idempotency_key": "restricted-1", "spec": {**DECK["spec"], "classification": "restricted"}}
    r = await arjun.post("/v1/pods/pod_alpha/exports", json=restricted)
    assert r.status_code == 403 and r.json()["reason"] == "classification_exceeds_clearance"
    assert (await alice.post("/v1/pods/pod_alpha/exports", json=restricted)).status_code == 201  # PM clearance


async def test_artifact_storage_failure_is_bounded(arjun, monkeypatch):
    import conftest
    from podos_api.objectstore import ObjectStoreError

    store = conftest.API_APP.state.services.store

    def broken_put(*a, **kw):
        raise ObjectStoreError("injected outage")

    monkeypatch.setattr(store, "put", broken_put)
    r = await arjun.post("/v1/pods/pod_alpha/exports", json={**DECK, "idempotency_key": "storage-fail-1"})
    assert r.status_code == 503 and r.json()["code"] == "artifact_storage_failed"
    assert (await arjun.get("/v1/pods/pod_alpha/artifacts")).json() == [], "no dangling artifact rows"
